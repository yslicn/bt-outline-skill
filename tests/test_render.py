#!/usr/bin/env python3
"""Render determinism tests: HTML structure, docx structure, same-source."""
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document

SKILL = Path(__file__).resolve().parent.parent
import sys
sys.path.insert(0, str(SKILL / "scripts"))
from render_html import render_html_bundle  # noqa: E402
from render_docx import render_docx_bundle  # noqa: E402


def sample_render_input() -> dict:
    return {
        "schema_version": "1.0", "stage_id": "stage_04_rendering_release", "project_name": "测试项目",
        "source_guide_full_sha256": "abc123",
        "meta": {
            "client": "测试客户", "transformation_theme": "营销服流程变革", "subtitle": "营销服流程变革 — 访谈提纲总览",
            "version": "v1.0", "date_text": "2026 年 8 月",
            "basis_meta": {"vs_count": 3, "l2_count": 4, "l3_count": 7, "interview_count": 2},
            "footer_text": "测试客户（eπ）营销服流程变革 — 访谈提纲总览",
        },
        "cover": {
            "eyebrow": "Management Consulting · Interview Guide",
            "h1": "测试客户（eπ）<br>营销服流程变革",
            "subtitle": "营销服流程变革 — 访谈提纲总览<br>3 大价值流 · 2 场专题访谈",
            "meta_lines": ["版本 v1.0 &nbsp;|&nbsp; 2026 年 8 月 &nbsp;|&nbsp; 内部工作文件"],
        },
        "background": ["本项目聚焦测试客户营销服流程变革。"],
        "methodology": {
            "progression": "五步法层层递进",
            "layer_definitions": [
                {"name": "破冰问题", "definition": "建立信任"},
                {"name": "唤醒问题", "definition": "核实挑战、机会与适用边界"},
            ],
            "capability_alignment": "问题标注VS与L2/L3",
        },
        "toc": [
            {"num": "01", "title": "品牌战略", "scope_line": "战略层 · 1项L2 · 1项L3", "target_audience": "集团高管", "filename": "01_品牌战略.html"},
            {"num": "02", "title": "媒介线索", "scope_line": "核心层 · 2项L2 · 5项L3", "target_audience": "投放团队", "filename": "02_媒介线索.html"},
        ],
        "vs_mapping_matrix": {
            "vs_codes": ["VS1", "VS2", "VS3"],
            "vs_short_names": ["品牌信任", "产品上市", "客户触达"],
            "rows": [
                {"interview_num": "01", "title": "品牌战略", "marks": [True, False, False]},
                {"interview_num": "02", "title": "媒介线索", "marks": [False, False, True]},
            ],
        },
        "guides": [
            {"num": "01", "title": "品牌战略", "vs_tags": ["VS1 品牌信任构建"], "layer_tag": "战略层 · 1项L2 · 1项L3",
             "meta": {"target_audience": "集团高管", "duration": "90 分钟", "capability_coverage_summary": "品牌战略（1项L3）"},
             "phases": [
                 {"phase_no": 1, "name": "暖场破冰", "layer_type": "intro", "title": "组织", "description": "破冰", "blocks": None, "questions": [{"text": "介绍团队？", "probes": [{"type": "追问", "text": "规模？"}]}]},
                 {"phase_no": 2, "name": "业务全景", "layer_type": "panorama", "title": "品牌全貌", "description": "全景", "blocks": [{"area_num": 1, "area_title": "品牌定位", "capability_ref": {"l2_id": "l2-01", "l2_name": "品牌战略", "l3_names": ["品牌定位"]}, "questions": [{"text": "品牌定位？", "probes": [{"type": "追问", "text": "客群？"}]}]}], "questions": None},
                 {"phase_no": 3, "name": "流程深挖", "layer_type": "deepdive", "title": "战略流程", "description": "深挖", "blocks": None, "questions": [{"text": "决策顺序？", "probes": [{"type": "追问", "text": "谁拍板？"}]}]},
                 {"phase_no": 4, "name": "挑战与机会探索", "layer_type": "awaken", "title": "挑战", "description": "唤醒", "blocks": None, "questions": [{"text": "目前客户对品牌的反馈主要有哪些？", "probes": [{"type": "引导思考", "text": "心智建设需多久？"}]}]},
                 {"phase_no": 5, "name": "变革期望", "layer_type": "expectation", "title": "诉求", "description": "期望", "blocks": None, "questions": [{"text": "最想改变什么？", "probes": [{"type": "追问", "text": "前提？"}]}]},
             ],
             "expected_output": ["品牌定位确认"]},
            {"num": "02", "title": "媒介线索", "vs_tags": ["VS3 客户触达"], "layer_tag": "核心层 · 2项L2 · 5项L3",
             "meta": {"target_audience": "投放团队", "duration": "90 分钟", "capability_coverage_summary": "媒介投放（1项L3）· 线索管理（4项L3）"},
             "phases": [
                 {"phase_no": 1, "name": "暖场破冰", "layer_type": "intro", "title": "组织", "description": "破冰", "blocks": None, "questions": [{"text": "同一团队？", "probes": [{"type": "追问", "text": "多少人？"}]}]},
                 {"phase_no": 2, "name": "业务全景", "layer_type": "panorama", "title": "投放全貌", "description": "全景", "blocks": [{"area_num": 1, "area_title": "投放", "capability_ref": {"l2_id": "l2-03", "l2_name": "媒介投放", "l3_names": ["投放管理"]}, "questions": [{"text": "投放策略？", "probes": [{"type": "追问", "text": "ROI？"}]}]}, {"area_num": 2, "area_title": "线索", "capability_ref": {"l2_id": "l2-04", "l2_name": "线索管理", "l3_count": 4}, "questions": [{"text": "线索链路？", "probes": [{"type": "追问", "text": "转化率？"}]}]}], "questions": None},
                 {"phase_no": 3, "name": "流程深挖", "layer_type": "deepdive", "title": "转化细节", "description": "深挖", "blocks": None, "questions": [{"text": "首次跟进时长？", "probes": [{"type": "追问", "text": "标准？"}]}]},
                 {"phase_no": 4, "name": "挑战与机会探索", "layer_type": "awaken", "title": "挑战", "description": "唤醒", "blocks": None, "questions": [{"text": "CPL水平？", "probes": [{"type": "引导思考", "text": "哪些环节运行顺畅？"}]}]},
                 {"phase_no": 5, "name": "变革期望", "layer_type": "expectation", "title": "诉求", "description": "期望", "blocks": None, "questions": [{"text": "升级什么？", "probes": [{"type": "追问", "text": "前提？"}]}]},
             ],
             "expected_output": ["投放全景", "线索链路"]},
        ],
    }


class TestRender(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.mkdtemp(prefix="gbt-render-")
        self.out = Path(self._tmp) / "candidate"

    def tearDown(self) -> None:
        shutil.rmtree(self._tmp, ignore_errors=True)

    def test_index_html_structure(self) -> None:
        ri = sample_render_input()
        render_html_bundle(ri, self.out, audience="interviewer")
        html = (self.out / "index.html").read_text(encoding="utf-8")
        for token in ("class=\"cover\"", "项目背景", "访谈设计方法论", "class=\"toc-grid\"", "class=\"toc-item\"", "价值流 × 访谈主题映射", "class=\"doc-footer\""):
            self.assertIn(token, html)
        self.assertEqual(html.count("class=\"toc-item\""), 2)

    def test_guide_html_structure(self) -> None:
        ri = sample_render_input()
        render_html_bundle(ri, self.out, audience="interviewer")
        html = (self.out / "01_品牌战略.html").read_text(encoding="utf-8")
        for token in ("class=\"header-bar\"", "class=\"interview-header\"", "class=\"interview-num\"", "tag tag-blue", "tag tag-teal", "class=\"meta-table\"", "class=\"output-box\"", "本场访谈预期输出"):
            self.assertIn(token, html)
        self.assertEqual(html.count("class=\"phase\""), 5)
        self.assertIn("引导思考", html)
        self.assertIn("追问", html)

    def test_docx_structure(self) -> None:
        ri = sample_render_input()
        render_docx_bundle(ri, self.out, audience="interviewer")
        doc = Document(self.out / "01_品牌战略.docx")
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        joined = "\n".join(texts)
        for token in ("访谈 01", "● Phase 1 · 暖场破冰", "● Phase 4 · 挑战与机会探索", "本场访谈预期输出"):
            self.assertIn(token, joined)
        self.assertGreaterEqual(len(doc.tables), 1)
        idx_doc = Document(self.out / "index.docx")
        toc_table = None
        for t in idx_doc.tables:
            if len(t.columns) >= 3 and t.rows and "编号" in t.rows[0].cells[0].text:
                toc_table = t
        self.assertIsNotNone(toc_table)
        self.assertEqual(len(toc_table.rows) - 1, 2)

    def test_same_source_bundle(self) -> None:
        ri = sample_render_input()
        render_html_bundle(ri, self.out, audience="interviewer")
        render_docx_bundle(ri, self.out, audience="interviewer")
        # HTML 与 docx 文件集一一对应（同源约束）
        html_nums = {p.name.split("_")[0] for p in self.out.glob("[0-9][0-9]_*.html")}
        docx_nums = {p.name.split("_")[0] for p in self.out.glob("[0-9][0-9]_*.docx")}
        self.assertEqual(html_nums, docx_nums)


if __name__ == "__main__":
    unittest.main()
