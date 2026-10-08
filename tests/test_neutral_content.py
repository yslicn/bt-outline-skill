"""Evidence, conditional-probe, audience and revision regression checks."""
from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from test_render import sample_render_input
from test_runtime import BTOutlineTestCase, run

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import ig_runtime as runtime
from render_docx import render_docx_bundle
from render_html import render_html_bundle


class TestNeutralContracts(BTOutlineTestCase):
    def test_no_negative_evidence_is_valid(self):
        baseline = self.baseline()
        baseline["strategic_tensions"] = []
        runtime.validate_stage01(baseline, {})
        runtime.validate_stage03(self.guide_full(), self.plan(), baseline)

    def test_new_source_must_resolve_and_verification_needs_note(self):
        baseline = self.baseline()
        item = baseline["strategic_tensions"][0]
        item["verification_status"] = "reported"
        with self.assertRaisesRegex(runtime.ContractError, "已登记材料"):
            runtime.validate_stage01(baseline, {"materials": [{"id": "another"}]})
        req = {"materials": [{"id": "strategy"}]}
        runtime.validate_stage01(baseline, req)
        item["verification_status"] = "verified"
        with self.assertRaisesRegex(runtime.ContractError, "verification_note"):
            runtime.validate_stage01(baseline, req)

    def test_question_ids_conditions_and_premise_refs(self):
        for field, value, message in [("id", "", "唯一稳定id"), ("premise_refs", ["missing"], "前提引用不存在")]:
            data = self.guide_full()
            data["guides"][0]["phases"][0]["questions"][0][field] = value
            with self.assertRaisesRegex(runtime.ContractError, message):
                runtime.validate_stage03(data, self.plan(), self.baseline())
        data = self.guide_full()
        probe = data["guides"][0]["phases"][3]["questions"][0]["probes"][0]
        probe["when"] = "confirmed"
        with self.assertRaisesRegex(runtime.ContractError, "具体condition"):
            runtime.validate_stage03(data, self.plan(), self.baseline())
        probe["condition"] = "受访者确认有影响客户反馈的具体情形"
        runtime.validate_stage03(data, self.plan(), self.baseline())

    def test_duplicate_question_ids_are_rejected(self):
        data = self.guide_full()
        data["guides"][1]["phases"][0]["questions"][0]["id"] = data["guides"][0]["phases"][0]["questions"][0]["id"]
        with self.assertRaisesRegex(runtime.ContractError, "唯一稳定id"):
            runtime.validate_stage03(data, self.plan(), self.baseline())

    def test_legacy_artifact_remains_readable(self):
        data = self.guide_full()
        data.pop("content_policy")
        for guide in data["guides"]:
            for phase in guide["phases"]:
                for q in (phase.get("questions") or []) + [q for b in phase.get("blocks") or [] for q in b["questions"]]:
                    q.pop("id")
                    for p in q["probes"]:
                        p.pop("when")
        runtime.validate_stage03(data, self.plan(), self.baseline())

    def test_client_and_interviewer_views_share_main_questions(self):
        ri = sample_render_input()
        question = ri["guides"][0]["phases"][3]["questions"][0]
        question["interviewer_context"] = "内部待核实假设"
        question["probes"][0].update(when="denied", text="哪些有效做法值得保留？")
        source = copy.deepcopy(ri)
        client, interviewer = self.ws / "client", self.ws / "interviewer"
        for directory, audience in [(client, "client"), (interviewer, "interviewer")]:
            render_html_bundle(ri, directory, audience=audience)
            render_docx_bundle(ri, directory, audience=audience)
        self.assertEqual(ri, source)
        html = (client / "01_品牌战略.html").read_text()
        internal = (interviewer / "01_品牌战略.html").read_text()
        doc = Document(client / "01_品牌战略.docx")
        client_text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        for token in ["VS1", "L2", "L3", "Phase", "能力域覆盖", "内部待核实假设", "哪些有效做法值得保留？"]:
            self.assertNotIn(token, html)
            self.assertNotIn(token, client_text)
        for token in ["VS1", "L2", "Phase", "内部待核实假设", "前提否定后"]:
            self.assertIn(token, internal)
        self.assertIn(question["text"], html)
        self.assertIn(question["text"], internal)
        for token in ["VS1", "L2", "L3", "能力域", "内部工作文件"]:
            self.assertNotIn(token, (client / "index.html").read_text())

    def test_public_context_and_answer_dimensions_survive_both_views(self):
        ri = sample_render_input()
        q = ri["guides"][0]["phases"][3]["questions"][0]
        q.update(context="以近期一笔订单的交期答复为例。", answer_hints=["可结合库存、已排任务与物料到位情况说明。"], interviewer_context="内部待核实判断")
        for audience in ["client", "interviewer"]:
            directory = self.ws / audience
            render_html_bundle(ri, directory, audience=audience)
            render_docx_bundle(ri, directory, audience=audience)
            html = (directory / "01_品牌战略.html").read_text()
            doc = Document(directory / "01_品牌战略.docx")
            text = "\n".join(p.text for p in doc.paragraphs)
            for value in [q["context"], *q["answer_hints"]]:
                self.assertIn(value, html)
                self.assertIn(value, text)
            if audience == "client":
                self.assertNotIn(q["interviewer_context"], html)
                self.assertNotIn(q["interviewer_context"], text)

    def test_source_vs_ids_with_gaps_are_not_renumbered(self):
        baseline = self.baseline()
        baseline["value_streams"][3]["id"] = "vs7"
        guides = self.guide_full()
        for g in guides["guides"]:
            for p in g["phases"]:
                for q in (p.get("questions") or []) + [q for b in p.get("blocks") or [] for q in b["questions"]]:
                    q["vs_refs"] = ["vs7" if ref == "vs4" else ref for ref in q["vs_refs"]]
        req = self.load_json("requirement.json")
        projected = runtime.project_render_input(guides, baseline, req, "hash")
        self.assertEqual(projected["vs_mapping_matrix"]["vs_codes"], ["VS1", "VS2", "VS3", "VS7"])

    def test_html_script_entry_point(self):
        import subprocess
        root = Path(__file__).resolve().parent.parent
        self.write_json("render_input.json", sample_render_input())
        result = subprocess.run([sys.executable, str(root / "scripts/render_html.py"), str(self.ws / "render_input.json"), str(self.ws / "html")], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.ws / "html/index.html").exists())


class TestRevisionClosure(BTOutlineTestCase):
    def open_issue(self):
        self.build_stages_01_03()
        # Begin a fresh content review, preserving stable question IDs.
        state = self.load_json("project_state.json")
        state["stages"][runtime.S03]["status"] = "in_progress"
        state["stages"][runtime.S03]["user_approval"] = None
        self.write_json("project_state.json", state)
        self.assertEqual(run("submit", str(self.ws), "--stage", runtime.S03).returncode, 0)
        question = self.guide_full()["guides"][0]["phases"][3]["questions"][0]
        self.write_json("issues.json", {"issues": [{"issue_id": "neutral-01", "question_id": question["id"], "status": "open", "problem": "追问未按否定回答继续", "required_fix": "增加适用的否定分支"}]})
        result = run("review", str(self.ws), "--stage", runtime.S03, "--role", "business_expert", "--decision", "REVISE", "--issues-file", str(self.ws / "issues.json"))
        self.assertEqual(result.returncode, 0, result.stderr)
        return question

    def test_cosmetic_revision_cannot_close_content_issue(self):
        question = self.open_issue()
        data = self.load_json(runtime.ARTIFACTS[runtime.S03])
        data["guides"][0]["meta"]["objectives"] = "只改显示元信息"
        self.write_json(runtime.ARTIFACTS[runtime.S03], data)
        self.assertEqual(run("submit", str(self.ws), "--stage", runtime.S03).returncode, 0)
        args = ("review", str(self.ws), "--stage", runtime.S03, "--role", "business_expert", "--decision", "PASS")
        result = run(*args)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("未关闭", result.stderr)
        self.write_json("closure.json", {"issues": [{"issue_id": "neutral-01", "question_id": question["id"], "status": "resolved", "resolution_note": "声称修复", "after": runtime.question_content(question)}]})
        result = run(*args, "--issues-file", str(self.ws / "closure.json"))
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("未修改", result.stderr)

    def test_probe_only_revision_can_be_verified(self):
        question = self.open_issue()
        data = self.load_json(runtime.ARTIFACTS[runtime.S03])
        changed = data["guides"][0]["phases"][3]["questions"][0]
        changed["probes"] = [{"type": "引导思考", "text": "哪些做法运行有效？", "when": "denied"}]
        self.write_json(runtime.ARTIFACTS[runtime.S03], data)
        self.assertEqual(run("submit", str(self.ws), "--stage", runtime.S03).returncode, 0)
        self.write_json("closure.json", {"issues": [{"issue_id": "neutral-01", "question_id": question["id"], "status": "resolved", "resolution_note": "按否定回答探索有效做法", "after": runtime.question_content(changed)}]})
        result = run("review", str(self.ws), "--stage", runtime.S03, "--role", "business_expert", "--decision", "PASS", "--issues-file", str(self.ws / "closure.json"))
        self.assertEqual(result.returncode, 0, result.stderr)
        ledger = self.load_json(runtime.S03 + "/review/issues_business_expert.json")
        self.assertEqual(ledger["issues"][0]["before"], runtime.question_content(question))
        self.assertEqual(ledger["issues"][0]["after"], runtime.question_content(changed))

    def test_closure_can_capture_visible_hint_change_without_retyping_after(self):
        q = self.open_issue()
        data = self.load_json(runtime.ARTIFACTS[runtime.S03])
        changed = data["guides"][0]["phases"][3]["questions"][0]
        changed["answer_hints"] = ["可介绍现有做法在哪些订单情形下运行有效。"]
        self.write_json(runtime.ARTIFACTS[runtime.S03], data)
        self.assertEqual(run("submit", str(self.ws), "--stage", runtime.S03).returncode, 0)
        self.write_json("closure.json", {"issues": [{"issue_id": "neutral-01", "question_id": q["id"], "status": "resolved", "resolution_note": "补充客户可见的有效做法回答范围"}]})
        result = run("review", str(self.ws), "--stage", runtime.S03, "--role", "business_expert", "--decision", "PASS", "--issues-file", str(self.ws / "closure.json"))
        self.assertEqual(result.returncode, 0, result.stderr)
        ledger = self.load_json(runtime.S03 + "/review/issues_business_expert.json")
        self.assertEqual(ledger["issues"][0]["after"], runtime.question_content(changed))
