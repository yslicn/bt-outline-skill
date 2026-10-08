#!/usr/bin/env python3
"""Render render_input.json to index.docx + NN_*.docx (same source as HTML)."""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path
from typing import Any
from render_views import prepare_view, probe_text

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

LATIN_FONT = "Aptos"


def detect_cjk_font() -> str:
    try:
        result = subprocess.run(["fc-list", ":lang=zh", "family"], capture_output=True, text=True, check=False, timeout=5)
        families = [line.split(",")[0].strip() for line in result.stdout.splitlines() if line.strip()]
        for preferred in ("PingFang SC", "Hiragino Sans GB", "Noto Sans CJK SC", "Arial Unicode MS"):
            if any(preferred.lower() in f.lower() for f in families):
                return preferred
        if families:
            return families[0]
    except (OSError, subprocess.SubprocessError):
        pass
    return "PingFang SC"


EAST_ASIA_FONT = detect_cjk_font()
PHASE_COLOR = {
    "intro": "8E6A00",
    "panorama": "007D79",
    "deepdive": "0043CE",
    "awaken": "DA1E28",
    "expectation": "24A148",
}
BLUE90 = "001D6C"


def _set_run_font(run: Any, *, size: float | None = None, bold: bool | None = None, italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = LATIN_FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{slot}"), LATIN_FONT if slot != "eastAsia" else EAST_ASIA_FONT)


def _write(paragraph: Any, value: Any, *, bold: bool = False, size: float | None = None, italic: bool = False, color: str | None = None) -> None:
    run = paragraph.add_run(str(value))
    _set_run_font(run, size=size, bold=bold, italic=italic, color=color)


def _cell(cell: Any, value: str, *, header: bool = False, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    _write(p, value, bold=bold or header, size=9)
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "80")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)
    if header:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "001D6C")
        tc_pr.append(shd)


def _plain_table(doc: Any, rows: int, cols: int) -> Any:
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    return table


def _para(doc: Any, value: str, *, bold: bool = False, size: float = 10, italic: bool = False, color: str | None = None, space_after: float = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    _write(p, value, bold=bold, size=size, italic=italic, color=color)
    return p


def _questions(doc: Any, questions: list[dict[str, Any]]) -> None:
    for q in questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Pt(12)
        _write(p, "Q ", bold=True, size=9.5, color=BLUE90)
        _write(p, q["text"], size=9.5)
        if q.get("interviewer_context"):
            _para(doc, "访谈员提示：" + q["interviewer_context"], size=9, color="6F6F6F")
        for probe in q.get("probes", []):
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(2)
            p2.paragraph_format.left_indent = Pt(24)
            _write(p2, "↳ ", bold=True, size=9, color="6F6F6F")
            text = probe_text(probe)
            _write(p2, text, size=9, color="6F6F6F")


def render_guide_docx(doc: Any, guide: dict[str, Any], ri: dict[str, Any]) -> None:
    meta = ri["meta"]
    internal = ri.get("audience") == "interviewer"
    bar = _plain_table(doc, 1, 1)
    _cell(bar.rows[0].cells[0], f'{meta["client"]} {meta["transformation_theme"]} · 访谈提纲（{guide["num"]} {guide["title"]}）', bold=True)
    _para(doc, f'访谈 {guide["num"]}', bold=True, size=10, color="0F62FE", space_after=1)
    _para(doc, guide["title"], bold=True, size=15, color=BLUE90, space_after=3)
    tags = "  ".join(guide.get("vs_tags", [])) + ("  |  " + guide["layer_tag"] if guide.get("layer_tag") else "")
    if tags:
        _para(doc, tags, size=8.5, color="6F6F6F", space_after=6)

    mt = _plain_table(doc, 3 if internal else 2, 2)
    m = guide["meta"]
    rows = [("访谈对象", m["target_audience"]), ("访谈时长", m["duration"])]
    if internal:
        rows.append(("能力域覆盖", m["capability_coverage_summary"]))
    for row, (k, v) in enumerate(rows):
        _cell(mt.rows[row].cells[0], k, header=True)
        _cell(mt.rows[row].cells[1], v)

    for phase in guide["phases"]:
        color = PHASE_COLOR[phase["layer_type"]]
        phase_label = f"● Phase {phase['phase_no']} · {phase['name']}" if internal else f"● {phase['name']}"
        _para(doc, phase_label, bold=True, size=11, color=color, space_after=1)
        if phase.get("title"):
            _para(doc, phase["title"], bold=True, size=10.5, color=BLUE90, space_after=1)
        if phase.get("description"):
            _para(doc, phase["description"], size=9, italic=True, color="6F6F6F", space_after=3)
        for block in phase.get("blocks") or []:
            cap = block.get("capability_ref") or {}
            ref = ""
            if cap:
                if cap.get("l3_names"):
                    ref = f"（L2：{cap['l2_name']} → {' · '.join(cap['l3_names'])}）"
                elif cap.get("l3_count"):
                    ref = f"（L2：{cap['l2_name']} → {cap['l3_count']}项L3）"
            _para(doc, f"{block['area_num']}. {block['area_title']}{ref}", bold=True, size=10, color=BLUE90, space_after=2)
            _questions(doc, block.get("questions") or [])
        _questions(doc, phase.get("questions") or [])

    _para(doc, "本场访谈预期输出", bold=True, size=10.5, color=BLUE90, space_after=2)
    for o in guide["expected_output"]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        _write(p, "•  ", size=9.5)
        _write(p, o, size=9.5)
    _para(doc, f'{meta["client"]} {meta["transformation_theme"]}  |  {meta["version"]}  |  {meta["date_text"]}', size=8, color="6F6F6F")


def render_index_docx(doc: Any, ri: dict[str, Any]) -> None:
    meta = ri["meta"]
    cover = ri["cover"]
    cover_tbl = _plain_table(doc, 1, 1)
    cell = cover_tbl.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    _write(p, cover["eyebrow"], size=9, bold=True, color="0F62FE")
    _write(p, "   " + cover["h1"].replace("<br>", " "), size=18, bold=True, color=BLUE90)
    for line in cover["subtitle"].split("<br>"):
        _para(doc, line, size=11, color="393939", space_after=2)
    for mline in cover["meta_lines"]:
        _para(doc, mline, size=9, color="6F6F6F", space_after=1)

    _para(doc, "项目背景", bold=True, size=13, color=BLUE90, space_after=2)
    for p_text in ri["background"]:
        _para(doc, p_text, size=9.5, space_after=3)
    _para(doc, "访谈设计方法论", bold=True, size=13, color=BLUE90, space_after=2)
    _para(doc, f"五层递进结构：{ri['methodology']['progression']}", size=9.5, space_after=2)
    for d in ri["methodology"]["layer_definitions"]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        _write(p, f"- {d['name']}：", size=9.5, bold=True)
        _write(p, d["definition"], size=9.5)
    alignment_label = "能力域对齐" if ri.get("audience") == "interviewer" else "业务主题覆盖"
    _para(doc, f"{alignment_label}：{ri['methodology']['capability_alignment']}", size=9.5, space_after=6)

    _para(doc, "访谈提纲清单", bold=True, size=13, color=BLUE90, space_after=3)
    toc_tbl = _plain_table(doc, len(ri["toc"]) + 1, 4)
    for col, head in enumerate(("编号", "标题", "范围", "对象")):
        _cell(toc_tbl.rows[0].cells[col], head, header=True)
    for i, item in enumerate(ri["toc"], start=1):
        _cell(toc_tbl.rows[i].cells[0], item["num"])
        _cell(toc_tbl.rows[i].cells[1], item["title"])
        _cell(toc_tbl.rows[i].cells[2], item["scope_line"])
        _cell(toc_tbl.rows[i].cells[3], item["target_audience"])
    _para(doc, "", size=6, space_after=4)

    matrix = ri["vs_mapping_matrix"]
    _para(doc, "价值流 × 访谈主题映射", bold=True, size=13, color=BLUE90, space_after=3)
    ncols = 1 + len(matrix["vs_codes"])
    m_tbl = _plain_table(doc, 2 + len(matrix["rows"]), ncols)
    _cell(m_tbl.rows[0].cells[0], "价值流", header=True)
    for j, code in enumerate(matrix["vs_codes"], start=1):
        _cell(m_tbl.rows[0].cells[j], code, header=True)
    _cell(m_tbl.rows[1].cells[0], "简称")
    for j, name in enumerate(matrix["vs_short_names"], start=1):
        _cell(m_tbl.rows[1].cells[j], name)
    for i, row in enumerate(matrix["rows"], start=2):
        _cell(m_tbl.rows[i].cells[0], f'{row["interview_num"]} {row["title"]}')
        for j, marked in enumerate(row["marks"], start=1):
            _cell(m_tbl.rows[i].cells[j], "●" if marked else "")

    _para(doc, "", size=6, space_after=4)
    _para(doc, f'{meta["client"]} {meta["transformation_theme"]} — 访谈提纲总览', bold=True, size=10, color=BLUE90, space_after=1)
    usage_label = "内部工作文件" if ri.get("audience") == "interviewer" else "访谈提纲"
    _para(doc, f'版本 {meta["version"]}  |  {meta["date_text"]}  |  {usage_label}', size=8.5, color="6F6F6F")


def render_docx_bundle(ri: dict[str, Any], out_dir: Path, *, audience: str = "client") -> list[Path]:
    ri = prepare_view(ri, audience)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    index_doc = Document()
    render_index_docx(index_doc, ri)
    index_path = out_dir / "index.docx"
    index_doc.save(index_path)
    paths.append(index_path)
    for guide in ri["guides"]:
        gdoc = Document()
        render_guide_docx(gdoc, guide, ri)
        gpath = out_dir / f"{guide['num']}_{guide['title']}.docx"
        gdoc.save(gpath)
        paths.append(gpath)
    return paths


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("render_input")
    parser.add_argument("out_dir")
    parser.add_argument("--audience", choices=["client", "interviewer"], default="client")
    args = parser.parse_args()
    ri = json.loads(Path(args.render_input).read_text(encoding="utf-8"))
    render_docx_bundle(ri, Path(args.out_dir), audience=args.audience)
    print(f"docx bundle written to {args.out_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
